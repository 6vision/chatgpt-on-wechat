"""Inline Word form values reach both document tools without changing the XML."""

from io import BytesIO

import pytest
import requests

from agent.tools.read.read import Read
from agent.tools.web_fetch.web_fetch import WebFetch
from common.office_text import iter_docx_body_text


@pytest.mark.parametrize("tool_name", ["read", "web_fetch"])
@pytest.mark.parametrize("placement", ["body", "cell"])
def test_inline_content_control_values_are_read(tmp_path, monkeypatch, tool_name, placement):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement

    def wrap(*runs):
        control = OxmlElement("w:sdt")
        properties = OxmlElement("w:sdtPr")
        alias = OxmlElement("w:alias")
        alias.set(docx.oxml.ns.qn("w:val"), "Control metadata")
        properties.append(alias)
        control.append(properties)
        content = OxmlElement("w:sdtContent")
        control.append(content)
        for run in runs:
            content.append(run)
        return control

    document = docx.Document()
    if placement == "body":
        paragraph = document.add_paragraph()
    else:
        paragraph = document.add_table(rows=1, cols=1).cell(0, 0).paragraphs[0]
    paragraph.add_run("Price: ")
    first = paragraph.add_run("123.")
    last = paragraph.add_run("45")
    paragraph._p.append(wrap(first._r, wrap(last._r)))
    paragraph.add_run("\tUSD\nApproved")
    path = tmp_path / "form.docx"
    document.save(path)

    if tool_name == "read":
        result = Read({"cwd": str(tmp_path)}).execute({"path": str(path)})
        text = result.result.get("content", "")
    else:
        tool = WebFetch({"cwd": str(tmp_path)})
        response = requests.Response()
        response.status_code = 200
        response.raw = BytesIO(path.read_bytes())
        response.headers["Content-Length"] = str(path.stat().st_size)
        monkeypatch.setenv("WEB_SECURITY_SSRF_PROTECTION", "false")
        monkeypatch.setattr(tool, "_safe_get", lambda *_args, **_kwargs: response)
        result = tool.execute({"url": "https://example.test/form.docx"})
        text = result.result

    assert result.status == "success", result.result
    assert "Price: 123.45\tUSD" in text
    assert "Approved" in text
    assert text.count("123.45") == 1
    assert "Control metadata" not in text

    loaded = docx.Document(path)
    original_xml = loaded.element.xml
    first_read = list(iter_docx_body_text(loaded))
    assert first_read == ["Price: 123.45\tUSD\nApproved"]
    assert list(iter_docx_body_text(loaded)) == first_read
    assert loaded.element.xml == original_xml


@pytest.mark.parametrize("has_content", [False, True])
def test_empty_inline_controls_preserve_surrounding_runs(has_content):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement

    document = docx.Document()
    paragraph = document.add_paragraph("Before ")
    control = OxmlElement("w:sdt")
    if has_content:
        control.append(OxmlElement("w:sdtContent"))
    paragraph._p.append(control)
    paragraph.add_run("after")

    assert list(iter_docx_body_text(document)) == ["Before after"]
