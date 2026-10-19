"""Document tools retain block-level Word content controls in reading order."""

from io import BytesIO

import pytest
import requests

from agent.tools.read.read import Read
from agent.tools.web_fetch.web_fetch import WebFetch
from common.office_text import iter_docx_body_text


@pytest.mark.parametrize("tool_name", ["read", "web_fetch"])
@pytest.mark.parametrize("placement", ["body", "cell"])
def test_block_content_controls_are_read(tmp_path, monkeypatch, tool_name, placement):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement

    def wrap(blocks):
        control = OxmlElement("w:sdt")
        properties = OxmlElement("w:sdtPr")
        alias = OxmlElement("w:alias")
        alias.set(docx.oxml.ns.qn("w:val"), "Control metadata")
        properties.append(alias)
        control.append(properties)
        content = OxmlElement("w:sdtContent")
        control.append(content)
        for block in blocks:
            content.append(block)
        return control

    document = docx.Document()
    if placement == "body":
        parent = document
    else:
        parent = document.add_table(rows=1, cols=1).cell(0, 0)

    parent.add_paragraph("Before control")
    first = parent.add_paragraph("Controlled paragraph")
    table = parent.add_table(rows=1, cols=1)
    table.cell(0, 0).text = "Controlled table"
    nested = parent.add_paragraph("Nested control")
    control = wrap([first._p, table._tbl, wrap([nested._p])])
    parent.add_paragraph("After control")._p.addprevious(control)
    path = tmp_path / "report.docx"
    document.save(path)

    if tool_name == "read":
        result = Read({"cwd": str(tmp_path)}).execute({"path": str(path)})
        text = result.result.get("content", "")
    else:
        tool = WebFetch({"cwd": str(tmp_path)})
        response = requests.Response()
        response.status_code = 200
        response.raw = BytesIO(path.read_bytes())
        response.headers["Content-Length"] = str(path.stat().st_size)
        monkeypatch.setenv("WEB_SECURITY_SSRF_PROTECTION", "false")
        monkeypatch.setattr(tool, "_safe_get", lambda *_args, **_kwargs: response)
        result = tool.execute({"url": "https://example.test/report.docx"})
        text = result.result

    assert result.status == "success", result.result
    expected = ["Before control", "Controlled paragraph", "Controlled table", "Nested control", "After control"]
    positions = [text.index(piece) for piece in expected]
    assert positions == sorted(positions)
    assert all(text.count(piece) == 1 for piece in expected)
    assert "Control metadata" not in text


@pytest.mark.parametrize("has_content", [False, True])
def test_empty_content_controls_preserve_surrounding_text(has_content):
    docx = pytest.importorskip("docx")
    from docx.oxml import OxmlElement

    document = docx.Document()
    document.add_paragraph("Before control")
    control = OxmlElement("w:sdt")
    if has_content:
        control.append(OxmlElement("w:sdtContent"))
    document.add_paragraph("After control")._p.addprevious(control)

    assert list(iter_docx_body_text(document)) == ["Before control", "After control"]
